from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from typing import List
import io
import PyPDF2
from docx import Document
from pptx import Presentation

# Initialize the new, separate FastAPI app
app = FastAPI(title="Examis Text Extractor API")

# Add CORS so the Flutter app is allowed to talk to this new URL
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/")
def read_root():
    return {"message": "Examis Text Extractor API is online!"}

@app.post("/extract-text")
async def extract_text(files: List[UploadFile] = File(...)):
    # Check if files were actually sent
    if not files:
        return {"success": False, "extracted_text": "", "error": "No files uploaded"}

    combined_text = ""

    # Loop through every file the Flutter app sent
    for file in files:
        filename = file.filename.lower()
        
        # Read the file directly into the server's RAM (Super fast!)
        contents = await file.read()
        file_bytes = io.BytesIO(contents)

        try:
            # 1. Handle PDF (Naturally ignores images)
            if filename.endswith('.pdf'):
                reader = PyPDF2.PdfReader(file_bytes)
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        combined_text += text + "\n"

            # 2. Handle DOCX (Only reads text paragraphs, ignores inline shapes)
            elif filename.endswith('.docx'):
                doc = Document(file_bytes)
                for para in doc.paragraphs:
                    if para.text.strip():
                        combined_text += para.text + "\n"

            # 3. Handle PPTX (Specifically checks for text frames to skip images)
            elif filename.endswith('.pptx'):
                prs = Presentation(file_bytes)
                for slide in prs.slides:
                    for shape in slide.shapes:
                        # THE IMAGE SKIPPER: If it has no text frame, skip it!
                        if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                            combined_text += shape.text + "\n"
            
            else:
                combined_text += f"\n[Skipped unsupported file format: {filename}]\n"

        except Exception as e:
            return {
                "success": False, 
                "extracted_text": "", 
                "error": f"Failed parsing {filename}: {str(e)}"
            }

    # Return the exact JSON structure your Flutter developer requested
    return {
        "success": True,
        "extracted_text": combined_text.strip(),
        "error": None
    }